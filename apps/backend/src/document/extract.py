"""In-memory extraction inspired by tuomin-open-source's block/Markdown flow.

Only extracted text is exported. Original containers and binary objects never
enter the output. Unsupported content blocks the entire conversion.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from io import BytesIO, StringIO
import re
import zipfile

from defusedxml import ElementTree as ET

FORMATS = ('txt', 'md', 'csv', 'json', 'yaml', 'yml', 'docx', 'xlsx', 'xlsm', 'xls', 'pdf')
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 32 * 1024 * 1024
MAX_TEXT_CHARS = 200_000
MAX_BLOCKS = 10_000
MAX_CELLS = 100_000
_W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


class DocumentError(Exception):
    """Static public error code, with no submitted text or parser diagnostics."""

    def __init__(self, code: str):
        self.code = code
        super().__init__(code)


@dataclass(frozen=True)
class Block:
    kind: str
    texts: tuple[str, ...]
    level: int = 0


@dataclass
class ExtractedDocument:
    blocks: list[Block] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    characters: int = 0
    cells: int = 0

    def add(self, kind: str, texts, level: int = 0):
        values = tuple(str(value) for value in texts)
        if kind == 'row':
            self.cells += len(values)
            if self.cells > MAX_CELLS or len(values) > 256:
                raise DocumentError('DOCUMENT_LIMIT_EXCEEDED')
        self.characters += sum(len(value) for value in values)
        if self.characters > MAX_TEXT_CHARS or len(self.blocks) >= MAX_BLOCKS:
            raise DocumentError('DOCUMENT_LIMIT_EXCEEDED')
        self.blocks.append(Block(kind, values, level))

    def warn(self, code: str):
        if code not in self.warnings:
            self.warnings.append(code)


def _package(data: bytes) -> zipfile.ZipFile:
    package = zipfile.ZipFile(BytesIO(data))
    try:
        items = package.infolist()
        names = [item.filename for item in items]
        if (len(items) > 2048 or len(names) != len(set(names))
                or sum(item.file_size for item in items) > MAX_EXPANDED_BYTES
                or any(item.flag_bits & 1 for item in items)
                or any(item.file_size > max(1, item.compress_size) * 200 for item in items)):
            raise DocumentError('DOCUMENT_LIMIT_EXCEEDED')
        # Validate every XML part, including the parts the text extractor omits.
        # defusedxml rejects DTD/entity expansion and external entities.
        for name in names:
            if name.lower().endswith(('.xml', '.rels', '.vml')):
                ET.fromstring(package.read(name), forbid_dtd=True)
        return package
    except BaseException:
        package.close()
        raise


def _word_text(node) -> str:
    pieces = []
    for child in node.iter():
        if child is not node and child.tag == _W + 'p' and pieces:
            pieces.append('\n')
        elif child.tag in {_W + 't', _W + 'delText', _W + 'instrText'}:
            pieces.append(child.text or '')
        elif child.tag == _W + 'tab':
            pieces.append('\t')
        elif child.tag in {_W + 'br', _W + 'cr'}:
            pieces.append('\n')
    return ''.join(pieces)


def _word_blocks(parent, output: ExtractedDocument):
    for node in parent:
        if node.tag == _W + 'p':
            text = _word_text(node)
            if not text.strip():
                continue
            style = node.find('./' + _W + 'pPr/' + _W + 'pStyle')
            heading = re.fullmatch(r'Heading([1-6])', style.get(_W + 'val', '')) if style is not None else None
            output.add('paragraph', (text,), int(heading[1]) if heading else 0)
        elif node.tag == _W + 'tbl':
            rows = list(node.iter(_W + 'tr'))
            if sum(len(list(row.iter(_W + 'tc'))) for row in rows) > MAX_CELLS:
                raise DocumentError('DOCUMENT_LIMIT_EXCEEDED')
            for row in rows:
                # Nested tables cannot be represented as a rectangular Markdown table.
                cells = list(row.findall(_W + 'tc'))
                if any(cell.find('.//' + _W + 'tbl') is not None for cell in cells):
                    raise DocumentError('DOCUMENT_UNSUPPORTED_CONTENT')
                output.add('row', tuple(_word_text(cell) for cell in cells))
            output.add('separator', ())
        else:
            _word_blocks(node, output)


def _docx(data: bytes, output: ExtractedDocument):
    with _package(data) as package:
        names = package.namelist()
        if 'word/document.xml' not in names:
            raise DocumentError('DOCUMENT_INVALID')
        if any(name.startswith(('word/embeddings/', 'word/media/', 'word/diagrams/'))
               or 'vbaProject' in name for name in names):
            raise DocumentError('DOCUMENT_UNSUPPORTED_CONTENT')
        parts = ['word/document.xml'] + sorted(
            name for name in names if re.fullmatch(r'word/(?:header\d*|footer\d*|footnotes|endnotes|comments)\.xml', name))
        for name in parts:
            root = ET.fromstring(package.read(name), forbid_dtd=True)
            if any(node.tag in {_W + 'altChunk', _W + 'object', _W + 'drawing', _W + 'pict'} for node in root.iter()):
                raise DocumentError('DOCUMENT_UNSUPPORTED_CONTENT')
            if any(node.tag in {_W + 'del', _W + 'ins', _W + 'moveFrom', _W + 'moveTo'} for node in root.iter()):
                output.warn('DOCX_REVISIONS_INCLUDED')
            if name != 'word/document.xml':
                output.warn('DOCX_ANCILLARY_INCLUDED')
            _word_blocks(root, output)


def _spreadsheet(data: bytes, output: ExtractedDocument):
    from openpyxl import load_workbook

    with _package(data) as package:
        names = package.namelist()
        if any(name.startswith(('xl/media/', 'xl/embeddings/', 'xl/charts/', 'xl/externalLinks/'))
               or name.startswith('xl/drawings/') and not re.fullmatch(r'xl/drawings/commentsDrawing\d+\.vml', name)
               for name in names):
            raise DocumentError('DOCUMENT_UNSUPPORTED_CONTENT')
        if any('vbaProject' in name for name in names):
            output.warn('SPREADSHEET_MACROS_OMITTED')
    book = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=False)
    try:
        cell_count = 0
        for sheet in book.worksheets:
            cell_count += sheet.max_row * sheet.max_column
            if cell_count > MAX_CELLS or sheet.max_column > 256:
                raise DocumentError('DOCUMENT_LIMIT_EXCEEDED')
            output.add('paragraph', (sheet.title,), 2)
            if (sheet.sheet_state != 'visible' or any(d.hidden for d in sheet.row_dimensions.values())
                    or any(d.hidden for d in sheet.column_dimensions.values())):
                output.warn('SPREADSHEET_HIDDEN_CONTENT_INCLUDED')
            for row in sheet.iter_rows():
                if not any(cell.value is not None or cell.comment is not None for cell in row):
                    continue
                values = []
                for cell in row:
                    # Formulas are text only: never evaluated, and never exported as XLSX/CSV.
                    if cell.data_type == 'f':
                        output.warn('SPREADSHEET_FORMULAS_AS_TEXT')
                    values.append('' if cell.value is None else str(cell.value))
                    if cell.comment is not None:
                        output.add('paragraph', (cell.comment.text, cell.comment.author or ''))
                        output.warn('SPREADSHEET_COMMENTS_INCLUDED')
                output.add('row', values)
            output.add('separator', ())
    finally:
        book.close()


def _xls(data: bytes, output: ExtractedDocument):
    import xlrd

    book = xlrd.open_workbook(file_contents=data, on_demand=True)
    try:
        cells = 0
        for sheet in book.sheets():
            cells += sheet.nrows * sheet.ncols
            if cells > MAX_CELLS or sheet.ncols > 256:
                raise DocumentError('DOCUMENT_LIMIT_EXCEEDED')
            output.add('paragraph', (sheet.name,), 2)
            for row in range(sheet.nrows):
                output.add('row', tuple(sheet.row_values(row)))
            output.add('separator', ())
        output.warn('XLS_CACHED_VALUES_ONLY')
    finally:
        book.release_resources()


def _pdf(data: bytes, output: ExtractedDocument):
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data), strict=True)
    if reader.is_encrypted:
        raise DocumentError('DOCUMENT_ENCRYPTED')
    if len(reader.pages) > 200:
        raise DocumentError('DOCUMENT_LIMIT_EXCEEDED')
    if reader.attachments:
        raise DocumentError('DOCUMENT_UNSUPPORTED_CONTENT')
    for page in reader.pages:
        resources = page.get('/Resources', {})
        objects = resources.get('/XObject', {})
        # Images and form XObjects require OCR / recursive layout analysis.
        if objects or page.get('/Annots'):
            raise DocumentError('DOCUMENT_UNSUPPORTED_CONTENT')
        contents = page.get_contents()
        if contents is not None and len(contents.get_data()) > MAX_EXPANDED_BYTES:
            raise DocumentError('DOCUMENT_LIMIT_EXCEEDED')
        text = page.extract_text() or ''
        if not text.strip() or '\ufffd' in text:
            raise DocumentError('DOCUMENT_NO_TEXT')
        output.add('paragraph', (text,))
    output.warn('PDF_LAYOUT_NOT_PRESERVED')


def extract_document(data: bytes, file_format: str) -> ExtractedDocument:
    if file_format not in FORMATS:
        raise DocumentError('DOCUMENT_FORMAT_UNSUPPORTED')
    if not data:
        raise DocumentError('DOCUMENT_NO_TEXT')
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError('DOCUMENT_LIMIT_EXCEEDED')
    output = ExtractedDocument()
    try:
        if file_format == 'docx':
            _docx(data, output)
        elif file_format in {'xlsx', 'xlsm'}:
            _spreadsheet(data, output)
        elif file_format == 'xls':
            _xls(data, output)
        elif file_format == 'pdf':
            _pdf(data, output)
        else:
            text = data.decode('utf-8-sig', errors='strict')
            if '\x00' in text:
                raise DocumentError('DOCUMENT_INVALID')
            if file_format == 'csv':
                cells = 0
                for row in csv.reader(StringIO(text), strict=True):
                    cells += len(row)
                    if len(row) > 256 or cells > MAX_CELLS:
                        raise DocumentError('DOCUMENT_LIMIT_EXCEEDED')
                    if row:
                        output.add('row', row)
            else:
                # Preserve complete text, including JSON keys, Markdown URLs and HTML.
                output.add('literal', (text,))
        if not output.blocks or not any(value.strip() for block in output.blocks for value in block.texts):
            raise DocumentError('DOCUMENT_NO_TEXT')
        return output
    except DocumentError:
        raise
    except Exception:
        raise DocumentError('DOCUMENT_INVALID') from None


def extraction_worker(data: bytes, file_format: str):
    """Bounded process entry point; no library diagnostics cross the boundary."""
    import logging
    import os
    import resource

    logging.disable(logging.CRITICAL)
    sink = os.open(os.devnull, os.O_WRONLY)
    try:
        os.dup2(sink, 1)
        os.dup2(sink, 2)
    finally:
        os.close(sink)
    resource.setrlimit(resource.RLIMIT_AS, (2 * 1024**3, 2 * 1024**3))
    try:
        return extract_document(data, file_format), None
    except DocumentError as error:
        return None, error.code
